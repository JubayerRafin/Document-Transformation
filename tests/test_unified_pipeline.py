from io import BytesIO
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile
import pymupdf
from transformation.office.package_ir import extract_package,write_package,inspect_package
from transformation.validation.privacy import detect
from transformation.pipeline import run_pipeline

HAS_OFFICE=bool(importlib.util.find_spec('docx') and importlib.util.find_spec('pptx'))


def office_fixture(kind):
    out=BytesIO()
    from PIL import Image
    image=BytesIO(); Image.new('RGB',(40,30),'#3768aa').save(image,format='PNG'); image.seek(0)
    if kind=='docx':
        from docx import Document
        doc=Document(); p=doc.add_paragraph(); p.add_run('Contact: '); p.add_run('maya@example.org').bold=True
        doc.add_paragraph('Owner: Maya Chen'); table=doc.add_table(rows=2,cols=2)
        table.cell(0,0).text='Item'; table.cell(0,1).text='Count'; table.cell(1,0).text='Printer'; table.cell(1,1).text='12'
        doc.sections[0].header.paragraphs[0].text='Internal document'
        doc.add_picture(image)
        doc.save(out)
    else:
        from pptx import Presentation
        from pptx.util import Inches
        deck=Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(1)).text='Contact: maya@example.org'
        slide.shapes.add_textbox(Inches(1),Inches(3),Inches(8),Inches(1)).text='Owner: Maya Chen'
        slide.shapes.add_picture(image,Inches(1),Inches(5),width=Inches(1))
        deck.save(out)
    return out.getvalue()


def pdf_fixture():
    with pymupdf.open() as pdf:
        page=pdf.new_page(); page.insert_text((40,60),'Contact: maya@example.org')
        page.insert_text((40,90),'Owner: Maya Chen')
        return pdf.tobytes()


class UnifiedTests(unittest.TestCase):
    def test_context_bridge_ignores_ambiguous_identical_labels(self):
        from extraction.context_bridge import build
        ir={'units':[{'regions':[{'cells':[['Field','Value'],['Title','Engineer']]},
                               {'cells':[['Field','Value'],['Title','A novel']]}]}]}
        context=build(ir)
        self.assertNotIn('Title',context)
        self.assertIn('Engineer',context)
        self.assertNotIn('A novel',context['Engineer'])

    def test_pii_rejects_hallucinated_source_spans(self):
        with patch('transformation.llm.ollama_native.chat',return_value=('{"findings":[{"value":"Invented Person","category":"Names"}]}',{})):
            with self.assertRaises(ValueError): detect([{'id':'a','text':'No person here.'}],['Names'],'unused','test')

    def test_rule_categories_do_not_require_model(self):
        with patch('transformation.llm.ollama_native.chat') as model:
            result=detect([{'id':'a','text':'Email maya@example.org; telephone +82 10 2345 6789'}],['Email addresses','Phone numbers'],'unused','test')
            self.assertEqual(len(result),2); model.assert_not_called()

    def test_no_operation_rejected(self):
        with self.assertRaises(ValueError): list(run_pipeline(b'','a.pdf',translate=False,redact=False))

    def test_visual_regions_are_validated(self):
        from transformation.analysis.document_analysis import analyze_pdf
        with patch('transformation.llm.ollama_native.chat',return_value=('{"purpose":"report","regions":[{"kind":"text","bbox":[0,0,2,1]}]}',{})):
            with self.assertRaises(ValueError): analyze_pdf(pdf_fixture(),'unused','test',True)

    @unittest.skipUnless(HAS_OFFICE,'Office export dependencies not installed in this interpreter')
    def test_all_nine_routes_with_real_writers_and_redaction(self):
        for source in ('pdf','docx','pptx'):
            data=pdf_fixture() if source=='pdf' else office_fixture(source)
            for output in ('pdf','docx','pptx'):
                with self.subTest(source=source,output=output),tempfile.TemporaryDirectory() as folder:
                    events=list(run_pipeline(data,'unrelated.'+source,output_kind=output,translate=False,
                        redact=True,pii_categories=['Email addresses'],pii_assist=False,table_strategy='native',root=Path(folder)))
                    last=events[-1]
                    self.assertEqual(last['status'],'generated',last.get('error'))
                    self.assertTrue(Path(last['source_ir']).is_file())
                    artifact=Path(last['output_file']).read_bytes()
                    if output=='pdf':
                        with pymupdf.open(stream=artifact,filetype='pdf') as pdf:
                            self.assertFalse(any(p.search_for('maya@example.org') for p in pdf))
                    else:
                        parsed=extract_package(artifact,'output.'+output)
                        texts=' '.join(r['text'] for u in parsed['units'] for r in u['regions'])
                        self.assertNotIn('maya@example.org',texts); self.assertIn('Maya Chen',texts)

    @unittest.skipUnless(HAS_OFFICE,'Office export dependencies not installed')
    def test_native_package_parts_and_namespaces_survive(self):
        data=office_fixture('docx'); ir=extract_package(data,'a.docx')
        region=next(r for u in ir['units'] for r in u['regions'] if r['text'].startswith('Contact:'))
        result,notes=write_package(data,'docx',{region['id']:['Contact: ','updated@example.org']})
        with zipfile.ZipFile(BytesIO(data)) as before,zipfile.ZipFile(BytesIO(result)) as after:
            self.assertEqual(before.namelist(),after.namelist())
            for name in before.namelist():
                if name!='word/document.xml': self.assertEqual(before.read(name),after.read(name))
            self.assertIn(b'xmlns:w14=',after.read('word/document.xml'))
        from docx import Document
        self.assertEqual(Document(BytesIO(result)).paragraphs[0].text,'Contact: updated@example.org')

    @unittest.skipUnless(HAS_OFFICE,'Office dependencies not installed')
    def test_cross_format_tables_use_translated_cells_once(self):
        from transformation.office.package_ir import refresh_tables
        from tools.legacy.cross_export import export
        ir=extract_package(office_fixture('docx'),'a.docx')
        for unit in ir['units']:
            for region in unit['regions']:
                if region.get('text')=='Printer': region['transformed_text']='Drucker'
        refresh_tables(ir)
        from docx import Document
        result=Document(BytesIO(export(ir,'docx')))
        self.assertEqual(result.tables[0].cell(1,0).text,'Drucker')
        self.assertNotIn('Printer',[p.text for p in result.paragraphs])

    @unittest.skipUnless(HAS_OFFICE,'Office dependencies not installed')
    def test_redaction_happens_before_translation(self):
        seen=[]
        def translate(endpoint,model,text,*args,**kwargs):
            seen.append(text); self.assertNotIn('maya@example.org',text)
            return text,{}
        def styled(endpoint,model,texts,*args,**kwargs):
            for text in texts: seen.append(text); self.assertNotIn('maya@example.org',text)
            return texts
        with tempfile.TemporaryDirectory() as folder,patch('transformation.office.styled_translation.translate_runs',side_effect=styled),patch('transformation.llm.ollama_native.translate_paragraph',side_effect=translate),patch('transformation.llm.translation_guard.audit_candidate',return_value={'faithful':True,'issues':[]}):
            last=list(run_pipeline(office_fixture('docx'),'form.docx',output_kind='docx',translate=True,
                redact=True,pii_categories=['Email addresses'],pii_assist=False,root=Path(folder)))[-1]
            self.assertEqual(last['status'],'generated',last.get('error'))
            self.assertTrue(any('[REDACTED]' in t for t in seen))

    @unittest.skipUnless(HAS_OFFICE,'Office dependencies not installed')
    def test_entire_protected_literal_skips_translation_and_audit(self):
        from docx import Document
        data=BytesIO(); doc=Document(); doc.add_paragraph('Maya Chen'); doc.save(data)
        with tempfile.TemporaryDirectory() as folder,patch('transformation.llm.ollama_native.chat') as chat:
            last=list(run_pipeline(data.getvalue(),'name.docx',output_kind='docx',translate=True,
                protected_values=['Maya Chen'],root=Path(folder)))[-1]
            self.assertEqual(last['status'],'generated',last.get('error'))
            chat.assert_not_called()


if __name__=='__main__': unittest.main()
