"""Build an allowlisted team snapshot. Run with the bundled document runtime."""
import hashlib
import json
from pathlib import Path
import shutil

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'outputs/team-pii-handoff-20261017'
PROJECT = OUT / 'TransformX_PII_Handoff/project'
BATCH = Path('runs/dataset-evaluation/translation-deep-doc-20261017-v1')

PAGES = [
('TransformX Code and PII Handoff', [
('p', 'Prepared for the TransformX teammate taking over PII implementation and evaluation. This snapshot documents the prototype as of 17 October 2026. Translation has run on six PDFs; the next task is to run Qwen 3.5 9B on the original PDFs, inspect the redacted outputs and evaluate detection and privacy separately.'),
('h', 'What is ready'),
('p', 'Streamlit exposes Translation only, PII only and combined modes. Deep Doc extraction integration now executes, including Docling tables and render and mask figure candidates. The shared coordinator supports a separate PII route. PII uses pattern detection for email and phone values and Qwen proposals for the remaining selected categories. The real six document PII GPU experiment and its human annotations are still pending.'),
('h', 'Current architecture'),
('code', 'Streamlit or batch CLI\n    -> document profile and extraction routing\n    -> Deep Doc adapters plus native geometry and PaddleOCR\n    -> TransformX content IR with regions and assets\n    -> Translation branch OR PII branch\n\nTranslation: TranslateGemma 12B through BabelDOC adapter\n             -> BabelDOC IR and PDF reconstruction\nPII:         patterns plus Qwen 3.5 9B exact span proposals\n             -> source PDF redaction and image processing\nBoth:        independent validation -> review -> evaluation'),
('h', 'Measured translation status'),
('table', [['Measure','Fallback','Deep Doc'],['Generated outputs','6 of 6','6 of 6'],['Table geometry F1 at IoU 0.50 and 0.75','1.000','1.000'],['Image placement F1 at both thresholds','0.966','0.966'],['Total processing time','885.7 seconds','1243.6 seconds']]),
('p', 'The layout measurements are almost identical. This is a before and after comparison, not a controlled timing experiment. New translation outputs have not inherited the fallback run human approvals or COMET scores. Generated is an operational status, not acceptance.')]),
('Code organization and ownership', [
('p', 'Start with transform_app.py, ui/transformx.py and transformation/pipeline.py. Production code is already grouped by responsibility. This handoff preserves imports and frozen runs; it adds navigation and portable packaging rather than relocating live modules during evaluation.'),
('table', [['Path','Responsibility'],['transform_app.py and ui/','Current Streamlit application and workflow controls'],['transformation/workflows.py','Independent run_translation and run_pii entry points'],['transformation/pipeline.py','Shared orchestration and PDF redact_pdf implementation'],['extraction/coordinator.py','Extraction routing and content IR conversion'],['extraction/deep_doc_adapter.py','Adapt prior Stage 1 methods and select guarded candidates'],['extraction/deep_doc_stage1/','Prior extraction algorithms retained with provenance'],['extraction/figures.py','Render and mask figure candidates and assets'],['transformation/ocr/ and tools/ocr/','Local or remote Paddle workers and OCR evidence'],['transformation/llm/','Ollama transport, translation adapters and protected text'],['transformation/validation/privacy.py','PII category policy and exact span detection'],['transformation/validation/privacy_verification.py','Fresh OCR checks on rendered output'],['transformation/pdf/ and layout/','BabelDOC integration, scan and fitting logic'],['transformation/office/','Native DOCX and PPTX routes requiring separate evaluation'],['tools/evaluation/','Batch registry, review bundles, scoring and comparisons'],['tools/gpu/ and tools/handoff/','Pod setup and portable team initialization'],['tests/','Regression tests and synthetic fixtures']]),
('p', 'app.py, main.py, extraction/pipeline.py and the evaluation/ folder also support the earlier extraction lab. They are retained for reproducibility. They are not the current Streamlit translation entry point. docs/TEAM_START.md and docs/ARCHITECTURE.md describe that earlier milestone; use this report for the current PII handoff.')]),
('Extraction and IR responsibilities', [
('p', 'The laptop profiles each page before detailed extraction. Searchable text uses native PDF evidence and Deep Doc pdfplumber adaptation. Docling supplies structural table candidates, with Camelot or pdfplumber fallbacks where appropriate. Native text coverage and geometry guards may reject candidates. Scans and selected raster figure content use PaddleOCR on the remote GPU. Render and mask candidates are heuristic and require inspection.'),
('p', 'There are two distinct representations. TransformX source-ir.json is the shared content IR for regions, coordinates, tables, assets and context. BabelDOC retains its own detailed PDF IR for glyph geometry, translation layout and reconstruction. The LLM receives selected text and bounded context, not both complete IR documents. The LLM does not draw the PDF.'),
('p', 'For PII only, source regions from TransformX IR go to the PII detector. The redactor applies findings to the source PDF. BabelDOC translation is not required for this PDF route. Scan and figure workers can redact raster content. Verification extracts or OCRs the resulting artifact again, so transformed-ir.json is output evidence and should not be confused with a second instruction document sent to the LLM.'),
('table', [['Artifact','Purpose'],['document-analysis.json','Page profile and routing evidence'],['deep-doc/extraction-evidence.json','Actual engines, accepted and rejected candidates'],['source-ir.json','Source text, region IDs, page boxes and assets'],['pii-findings.json','Detected literals, categories, offsets and methods'],['redaction-audit.json','How each source finding was located and removed'],['transformed-ir.json','Content extracted from the generated artifact'],['independent-validation.json','Residual values, geometry and uninspected content'],['privacy-verification.json','Fresh rendered output OCR checks when invoked'],['output/result.pdf and manifest.json','Final artifact, settings, status and limitations']]),
('p', 'In the new translation run, Docling tables were accepted in Samples02 through 04. Sample01 retains one uncertain OCR region, and one table candidate was rejected by the coverage guard. These are unresolved findings, not conditions to suppress in evaluation.')]),
('Install the local environments', [
('p', 'Extract the ZIP into a short folder, such as C:\\TransformXTeam. Open the project subfolder in VS Code. Run all following PowerShell commands from that project folder. Use Python 3.11 for local environments; the GPU installer uses the pod Python. Keep the BabelDOC environment separate to avoid dependency conflicts.'),
('code', 'py -3.11 -m venv .venv-paddleocr\n.\\.venv-paddleocr\\Scripts\\python.exe -m pip install -r requirements.txt -r requirements-deep-doc.txt\n\npy -3.11 -m venv .venv-babeldoc\n.\\.venv-babeldoc\\Scripts\\python.exe -m pip install -r requirements-transformation.txt\n\npy -3.11 -m venv .venv-evaluation\n.\\.venv-evaluation\\Scripts\\python.exe -m pip install -r requirements.txt -r requirements-evaluation.txt'),
('p', 'The app environment needs extraction packages, while the evaluation environment can read saved evidence and compute detailed metrics. COMET is optional for translation and is not a PII metric. Model weights, virtual environments, credentials and third party caches are deliberately excluded from the ZIP. Native Office export may also need LibreOffice and is outside the six PDF PII experiment.'),
('h', 'Initialize the portable experiment once'),
('code', '.\\.venv-paddleocr\\Scripts\\python.exe -m tools.handoff.initialize_pii'),
('p', 'The initializer verifies the six source hashes, relocates path fields in the copied run evidence and writes configs/pii-team.local.json. It records the original experiment fingerprint and calculates a local fingerprint for the relocated policy. It does not rerun translation or transfer human approvals. The untouched original index and manifest copies remain under evidence/.'),
('p', 'The completed translation index is included because the current batch runner requires a completed translation phase before its PII phase. This is an experiment sequencing guard, not a dependency on translated PDFs: PII starts again from the original sources. Do not run prepare on an initialized package because it can replace the registered source metadata. Do not change policy or runtime source code after creating PII annotations; use a new experiment root if the code changes.')]),
('Connect the GPU and run PII', [
('p', 'Use your own pod address, SSH port and private key. Do not reuse Jubayer connection details. Verify nvidia-smi and HTTPS downloads inside the pod before installation. Upload tools/gpu/setup_transformx.sh through scp, then run the following Linux command inside the root prompt.'),
('code', 'bash /workspace/setup_transformx.sh pii'),
('p', 'Wait for Setup finished. The installer checks CUDA tensor placement and convolution with numerical tolerance, installs PaddleOCR, and caches the detector and recognition models. A successful pip installation alone is not a successful OCR setup. The pii option pulls Qwen; all also pulls TranslateGemma.'),
('p', 'In a separate Windows PowerShell terminal, open the tunnel and keep it running. Replace each placeholder with your actual value.'),
('code', 'ssh.exe -N -o ExitOnForwardFailure=yes -o ServerAliveInterval=30 -p SSH_PORT -i "PRIVATE_KEY_PATH" -L 127.0.0.1:11435:127.0.0.1:11434 root@GPU_IP'),
('p', 'Edit configs/pii-team.local.json: endpoint stays http://127.0.0.1:11435/v1; pii_model stays qwen3.5:9b; fill ocr_remote.host, port and key. Keep table_strategy set to docling for this experiment. All six supported categories are selected. Keep visual analysis off, manual protected values empty and context empty for the registered baseline. Protected values are translation preservation settings, not PII exemptions.'),
('code', '.\\.venv-paddleocr\\Scripts\\python.exe -m tools.evaluation.dataset preflight --config configs/pii-team.local.json --phase pii\n.\\.venv-paddleocr\\Scripts\\python.exe -m tools.evaluation.dataset run --config configs/pii-team.local.json --phase pii'),
('p', 'For an interrupted unchanged experiment add --resume to the run command. Inspect the first source and redacted output as soon as it finishes. The runner checks model VRAM evidence; a failure before model inference can produce a secondary GPU verification warning. Read the original error and the nested OCR detection or recognition log. Stop the paid GPU when inference is complete; retained storage may still be billed.')]),
('How to annotate PII ground truth', [
('p', 'After PII execution, review templates are written to ground-truth/Sample01.pii.json through Sample06.pii.json under the initialized batch. review-bundles/SampleXX/pii contains rendered source and output pages plus regions.csv. Use the PII source IR from that run; extraction and region IDs can differ from the translation run. Never copy translation region IDs or the model findings into the gold labels without checking the original page.'),
('p', 'For every region, compare extracted_text with the original PDF, enter an independently checked source_reference and set source_reviewed=true only when verified. Put each sensitive occurrence into the region pii array. Offsets are zero based Python character positions, with start included and end excluded. They refer to extracted_text, not the corrected source_reference.'),
('code', 'Example extracted_text: "Contact Maya"\n\n"pii": [\n  {"start": 8, "end": 12, "value": "Maya",\n   "category": "Names"}\n]\n\nRequired invariant:\nextracted_text[start:end] == value'),
('table', [['Category string','Annotation scope'],['Email addresses','Visible email occurrences'],['Phone numbers','Phone and extension values under the agreed policy'],['Names','Person names under the agreed policy'],['Addresses','Personal location or address spans'],['Identifiers','Sensitive IDs under the agreed policy'],['Bank accounts','Account values under the agreed policy']]),
('p', 'Before labeling, agree whether employee IDs, salary, signatures, company addresses and document numbers are in scope. The current six categories are not a complete organizational privacy policy. Record ambiguous cases and apply one consistent decision across all PDFs. Mark empty pii arrays as reviewed negatives only after checking the whole region.'),
('p', 'Put PII visible on the page but absent from extracted text into missing_source_regions, including page, box, visible text and category. Record visible residuals in privacy_residuals and unnecessary masking in over_redactions. Set pii_ground_truth_reviewed, source_coverage_reviewed, visual_reviewed and privacy_reviewed only after their respective checks are complete. Do not set human_accepted merely because the output opens.')]),
('Evaluate detection and actual removal', [
('p', 'Run the detailed evaluator after filling all six PII ground truth files. Do not run COMET for PII; it evaluates translation meaning rather than entity detection or removal.'),
('code', '.\\.venv-evaluation\\Scripts\\python.exe -m tools.evaluation.detailed evaluate --config configs/pii-team.local.json --phase pii'),
('p', 'Read pii/detailed/report.html, report.json and summary.csv under the batch folder. These are the authoritative region based PII metrics. The simpler dataset report uses a separate annotations.json schema; do not assume populating ground-truth automatically populates that older form. Prefer the detailed report for the teammate deliverable.'),
('table', [['Metric or check','Interpretation'],['Exact span precision','TP / (TP + FP); correct region, category and exact offsets'],['Exact span recall','TP / (TP + FN); misses on annotated extracted text'],['Exact span F1','2PR / (P + R); report counts and scores per document'],['Overlap span score','One to one overlapping spans in the same region and category'],['Per category scores','Separate the six classes so weak classes are visible'],['CER and WER','Reviewed source transcription versus extracted text; OCR quality'],['Privacy residual review','PII remaining in visible pixels, text or hidden content'],['Over redaction review','Useful non sensitive content unnecessarily removed'],['Layout preservation','Page geometry, figures, clipped text and affected tables']]),
('p', 'Report macro and micro aggregates only if explicitly implemented and verified. Existing detailed reports provide per document and per category scores; do not invent a pooled percentage. True negatives are not defined for arbitrary text spans, so a generic accuracy percentage is misleading. Show TP, FP and FN. With no gold or predicted spans, inspect denominators and label a reviewed negative explicitly.'),
('p', 'Detection F1 measures only extracted regions. It can be high even when OCR missed a sensitive address in an image. Actual removal must therefore be checked independently against original pages, rendered redacted pages, searchable text, metadata and attachments. Fresh rendered OCR is useful but can also miss text. No automated check in this snapshot certifies complete anonymization.')]),
('Remaining work and completion criteria', [
('table', [['Priority','Next work'],['1','Verify clean GPU setup and run one PDF end to end in PII only mode'],['2','Run all six original PDFs and retain failures, settings and Qwen digest'],['3','Independently label sensitive spans and original page coverage'],['4','Evaluate exact and overlap scores by document and category'],['5','Review raster residuals, hidden content and over redaction'],['6','Fix generic detector, mapping or removal defects with regression tests'],['7','Freeze a new version and rerun affected PDFs before accepting it'],['8','Expand to unseen scans, multilingual PII, complex tables and Office files']]),
('h', 'Known limitations to investigate'),
('p', 'The six PDF PII model experiment is unvalidated. Email and phone patterns are limited, and semantic names or IDs may require more context. Scans currently redact an entire detected line, which can remove useful content. Figure candidates can overlap native text and remain blocked. Sample01 OCR is uncertain; Sample04 has image and layout issues in the translation run. Those translation findings are risk examples, not automatically verified PII defects.'),
('p', 'Check the verification language in PII only mode: privacy_verification.py currently passes options.target to fresh OCR even though the unchanged source language may be English. Also audit child figure findings: some are saved under nested figure runs and should not be silently omitted from end to end accounting. Treat the manifest field pii_detection_complete=false as a stated completeness limitation, not a reason to set it true manually.'),
('h', 'Regression checks before model testing'),
('code', '.\\.venv-paddleocr\\Scripts\\python.exe -m unittest tests.test_separate_workflows tests.test_detailed_evaluation tests.test_remote_ocr_transport tests.test_handoff -v'),
('p', 'These local tests cover operation routing, synthetic redaction, annotation provenance and metric behavior. They do not replace real Qwen GPU inference or a reviewed privacy dataset. Do not hardcode sample names, expected entities or corrected PDF text into production. References and human labels remain evaluation inputs only.'),
('h', 'What to return to the team'),
('p', 'Submit source changes, a versioned configuration without credentials, source hashes, model digest, package versions, device evidence, six output paths, reviewed annotations, detailed JSON and CSV metrics, screenshots of failures, known limitations and a rerun command. Preserve this snapshot and previous reports. Agree acceptance thresholds with the mentor before claiming the generalized prototype is privacy safe.')])
]


def docx_report(path):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.7)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles['Normal'].font.size = Pt(10.5)
    doc.styles['Normal'].paragraph_format.space_after = Pt(6)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.08
    doc.styles['Title'].font.size = Pt(22)
    doc.styles['Heading 1'].font.size = Pt(17)
    doc.styles['Heading 2'].font.size = Pt(12)
    foot = section.footer.paragraphs[0]
    foot.text = 'TransformX  |  PII teammate handoff  |  17 October 2026'
    foot.runs[0].font.size = Pt(9)
    for n, (title, blocks) in enumerate(PAGES):
        if n:
            doc.add_page_break()
        doc.add_paragraph(title, 'Title' if n == 0 else 'Heading 1')
        for kind, content in blocks:
            if kind == 'p':
                doc.add_paragraph(content)
            elif kind == 'h':
                doc.add_paragraph(content, 'Heading 2')
            elif kind == 'code':
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(7)
                p.paragraph_format.line_spacing = 1
                r = p.add_run(content)
                r.font.name = 'Consolas'; r.font.size = Pt(8.2)
            elif kind == 'table':
                table = doc.add_table(rows=0, cols=len(content[0]))
                table.style = 'Table Grid'
                for index, row in enumerate(content):
                    cells = table.add_row().cells
                    for cell, text in zip(cells, row):
                        cell.text = text
                        for p in cell.paragraphs:
                            p.paragraph_format.space_after = Pt(4)
                            p.paragraph_format.space_before = Pt(4)
                            for run in p.runs:
                                run.font.size = Pt(9)
                                run.bold = index == 0
                        if index == 0:
                            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'DCE6F1')
                            cell._tc.get_or_add_tcPr().append(shade)
                    props = table.rows[-1]._tr.get_or_add_trPr()
                    props.append(OxmlElement('w:cantSplit'))
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
    doc.save(path)


def main():
    if (PROJECT / 'configs/pii-team.local.json').exists():
        raise ValueError('Refusing to rebuild an initialized teammate workspace.')
    PROJECT.mkdir(parents=True, exist_ok=True)
    # Explicit allowlist. No virtualenvs, secrets, SSH keys, old outputs or model caches.
    for folder in ('transformation', 'extraction', 'ui', 'tools', 'tests', 'evaluation'):
        for src in (ROOT / folder).rglob('*.py'):
            if '__pycache__' in src.parts:
                continue
            dest = PROJECT / src.relative_to(ROOT)
            dest.parent.mkdir(parents=True, exist_ok=True); shutil.copy2(src, dest)
    for name in ('transform_app.py', 'app.py', 'main.py', 'babeldoc_adapter.py', 'README.md', 'LICENSE', 'Start TransformX.cmd', 'Start Lab.cmd'):
        shutil.copy2(ROOT / name, PROJECT / name)
    for src in ROOT.glob('requirements*.txt'):
        shutil.copy2(src, PROJECT / src.name)
    for folder in ('config', 'docs'):
        shutil.copytree(ROOT / folder, PROJECT / folder, dirs_exist_ok=True)
    (PROJECT / 'configs').mkdir(exist_ok=True)
    shutil.copy2(ROOT / 'configs/gpu-evaluation.example.json', PROJECT / 'configs/gpu-evaluation.example.json')
    setup = ROOT / 'tools/gpu/setup_transformx.sh'
    (PROJECT / setup.relative_to(ROOT)).parent.mkdir(parents=True, exist_ok=True)
    (PROJECT / setup.relative_to(ROOT)).write_bytes(setup.read_bytes())
    config = json.loads((ROOT / 'configs/evaluation-deep-doc-20261017.local.json').read_text())
    source_map = {}
    for item in config['documents']:
        relative = 'samples/' + item['id'] + '.pdf'
        (PROJECT / 'samples').mkdir(exist_ok=True)
        shutil.copy2(item['path'], PROJECT / relative)
        source_map[item['path']] = relative
    index = json.loads((ROOT / BATCH / 'translation/index.json').read_text())
    # Keep completed runs only; omit bulky typesetting work PDFs and duplicated archives.
    completed = {Path(r['manifest']).parent.name for r in index['records']}
    for src in (ROOT / BATCH).rglob('*'):
        if not src.is_file():
            continue
        rel = src.relative_to(ROOT / BATCH)
        if rel.parts[0] in ('ground-truth', 'review-bundles'):
            continue
        if rel.parts[0] == 'shared-ocr':
            continue  # Teammate gets fresh OCR evidence; old shared cache is not portable.
        if 'work' in rel.parts or src.suffix == '.zip':
            continue
        if len(rel.parts) >= 3 and rel.parts[:2] == ('translation', 'runs') and rel.parts[2] not in completed:
            continue
        dest = PROJECT / BATCH / rel
        dest.parent.mkdir(parents=True, exist_ok=True); shutil.copy2(src, dest)
    evidence = PROJECT / 'evidence'; evidence.mkdir(exist_ok=True)
    shutil.copy2(ROOT / BATCH / 'translation/index.json', evidence / 'original-translation-index.json')
    for record in index['records']:
        shutil.copy2(record['manifest'], evidence / (record['document_id'] + '-original-manifest.json'))
    baseline = ROOT / 'runs/dataset-evaluation/translation-baseline-20261015-v2'
    for rel in ('translation/detailed/summary.csv', 'layout-standard-metrics/report.json'):
        dest = evidence / ('fallback-' + Path(rel).name)
        shutil.copy2(baseline / rel, dest)
    config['ocr_remote'] = {'host': '', 'port': 22, 'key': 'SET_YOUR_OWN_PRIVATE_KEY_PATH'}
    config['babeldoc_python'] = '.venv-babeldoc/Scripts/python.exe'
    for item in config['documents']:
        item['path'] = 'samples/' + item['id'] + '.pdf'
    meta = {'original_project_root': str(ROOT), 'batch_relative': BATCH.as_posix(),
            'source_map': source_map, 'configuration': config,
            'omitted': ['virtual environments', 'SSH private keys and tokens', 'model weights',
                        'shared OCR cache', 'BabelDOC work PDFs', 'old failed runs', 'output-specific human approvals']}
    (PROJECT / 'handoff-package.json').write_text(json.dumps(meta, indent=2), encoding='utf-8')
    report = OUT / 'TransformX_PII_Handoff/TransformX_Code_and_PII_Handoff.pdf'
    text = '\n\n'.join(title + '\n' + '\n\n'.join(('\n'.join(' | '.join(row) for row in content) if kind == 'table' else content) for kind, content in blocks) for title, blocks in PAGES)
    (PROJECT / 'docs/PII_HANDOFF.txt').write_text(text, encoding='utf-8')
    code_map = '\n\n'.join(('\n'.join(' | '.join(row) for row in c) if k == 'table' else c) for k, c in PAGES[1][1])
    (PROJECT / 'docs/CODE_MAP.txt').write_text(code_map, encoding='utf-8')
    shutil.copy2(PROJECT / 'docs/PII_HANDOFF.txt', ROOT / 'docs/PII_HANDOFF.txt')
    shutil.copy2(PROJECT / 'docs/CODE_MAP.txt', ROOT / 'docs/CODE_MAP.txt')
    (OUT / 'TransformX_PII_Handoff/START_HERE.txt').write_text(
        'TransformX PII teammate handoff\n\n1. Read TransformX_Code_and_PII_Handoff.pdf.\n'
        '2. Open project/ in VS Code. Use the three local environments described in the report.\n'
        '3. Run python -m tools.handoff.initialize_pii with the project environment.\n'
        '4. Set your own GPU SSH details in configs/pii-team.local.json. Never copy a private key into this ZIP.\n'
        '5. Run the PII phase on the original six PDFs, then annotate and evaluate.\n\n'
        'No real six PDF PII model evaluation or human approval is claimed.\n'
        'Saved translation outputs and original sources are included for continuity.\n'
        'No personal environments, model weights or SSH credentials are included.\n', encoding='utf-8')
    from tools.handoff.build_report_pdf import main as build_pdf
    build_pdf()
    inventory = [{'path': p.relative_to(PROJECT.parent).as_posix(), 'bytes': p.stat().st_size,
                  'sha256': hashlib.sha256(p.read_bytes()).hexdigest()}
                 for p in PROJECT.parent.rglob('*') if p.is_file() and p.suffix != '.docx' and p.name != 'PACKAGE_INVENTORY.json']
    (PROJECT.parent / 'PACKAGE_INVENTORY.json').write_text(json.dumps(inventory, indent=2), encoding='utf-8')
    print(report)
    print('Packaged files:', len(inventory), 'uncompressed MB:', round(sum(x['bytes'] for x in inventory)/1e6, 1))


if __name__ == '__main__':
    main()
