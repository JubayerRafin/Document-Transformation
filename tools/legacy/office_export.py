"""Visual Office exports. Pages are images, explicitly not editable reconstruction."""
from io import BytesIO
import pymupdf


def export_visual(pdf_path, kind):
    output=BytesIO()
    with pymupdf.open(pdf_path) as pdf:
        if kind=='pptx':
            from pptx import Presentation
            from pptx.util import Pt
            deck=Presentation()
            deck.slide_width=Pt(pdf[0].rect.width)
            deck.slide_height=Pt(pdf[0].rect.height)
            for page in pdf:
                slide=deck.slides.add_slide(deck.slide_layouts[6])
                scale=min(deck.slide_width/page.rect.width,deck.slide_height/page.rect.height)
                width,height=int(page.rect.width*scale),int(page.rect.height*scale)
                image=BytesIO(page.get_pixmap(matrix=pymupdf.Matrix(2,2),alpha=False).tobytes('png'))
                slide.shapes.add_picture(image,int((deck.slide_width-width)/2),int((deck.slide_height-height)/2),width=width,height=height)
            deck.save(output)
        elif kind=='docx':
            from docx import Document
            from docx.shared import Pt
            from docx.enum.section import WD_SECTION_START
            document=Document()
            for i,page in enumerate(pdf):
                section=document.sections[0] if i==0 else document.add_section(WD_SECTION_START.NEW_PAGE)
                section.page_width=Pt(page.rect.width)
                section.page_height=Pt(page.rect.height)
                section.top_margin=section.bottom_margin=Pt(2)
                section.left_margin=section.right_margin=Pt(2)
                paragraph=document.add_paragraph()
                paragraph.paragraph_format.space_before=Pt(0)
                paragraph.paragraph_format.space_after=Pt(0)
                paragraph.paragraph_format.line_spacing=1
                image=BytesIO(page.get_pixmap(matrix=pymupdf.Matrix(2,2),alpha=False).tobytes('png'))
                # Reserve room for the paragraph baseline; preserve aspect ratio.
                scale=min((page.rect.width-4)/page.rect.width,(page.rect.height-20)/page.rect.height)
                paragraph.add_run().add_picture(image,width=Pt(page.rect.width*scale),height=Pt(page.rect.height*scale))
            document.save(output)
        else:
            raise ValueError('Choose docx or pptx.')
    return output.getvalue()
