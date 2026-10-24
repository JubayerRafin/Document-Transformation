"""Explicitly reflowed, editable Office exports from the shared content IR."""
from io import BytesIO


def blocks(ir):
    for unit in ir['units']:
        for region in unit['regions']:
            if region.get('cells'): yield 'table',region['cells']
            elif region.get('text','').strip() and not region.get('metadata',{}).get('inside_table'):
                yield 'text',region.get('transformed_text',region['text'])


def export(ir,kind,images=()):
    output=BytesIO()
    if kind=='docx':
        from docx import Document
        doc=Document()
        for role,value in blocks(ir):
            if role=='text': doc.add_paragraph(value)
            else:
                table=doc.add_table(rows=len(value),cols=max(map(len,value)))
                table.style='Table Grid'
                for i,row in enumerate(value):
                    for j,cell in enumerate(row): table.cell(i,j).text=str(cell or '')
        for image in images:
            from docx.shared import Inches
            doc.add_paragraph().add_run().add_picture(BytesIO(image),width=Inches(6))
        doc.save(output)
    elif kind=='pptx':
        from pptx import Presentation
        from pptx.util import Inches,Pt
        deck=Presentation(); deck.slide_width=Inches(13.333); deck.slide_height=Inches(7.5)
        # Bounded content per slide; no assumption that a Word story is a page.
        for role,value in blocks(ir):
            if role=='table':
                for start in range(0,len(value),10):
                    rows=value[start:start+10]; slide=deck.slides.add_slide(deck.slide_layouts[6])
                    table=slide.shapes.add_table(len(rows),max(map(len,rows)),Inches(.5),Inches(.5),Inches(12.3),Inches(6.3)).table
                    for i,row in enumerate(rows):
                        for j,cell in enumerate(row): table.cell(i,j).text=str(cell or '')
            else:
                chunks=[value[i:i+650] for i in range(0,len(value),650)]
                for chunk in chunks:
                    slide=deck.slides.add_slide(deck.slide_layouts[6])
                    box=slide.shapes.add_textbox(Inches(.6),Inches(.6),Inches(12),Inches(6.1))
                    box.text_frame.word_wrap=True; box.text_frame.text=chunk
                    for p in box.text_frame.paragraphs: p.font.size=Pt(22)
        for image in images:
            from PIL import Image
            picture=Image.open(BytesIO(image)); w,h=picture.size
            scale=min(Inches(12)/w,Inches(6.2)/h)
            slide=deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_picture(BytesIO(image),int((deck.slide_width-w*scale)/2),int((deck.slide_height-h*scale)/2),width=int(w*scale),height=int(h*scale))
        if not len(deck.slides): deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(output)
    else: raise ValueError('Editable reflow export supports DOCX and PPTX.')
    return output.getvalue()


def pdf_reflow(ir,images=()):
    import html
    import pymupdf
    import tempfile
    from pathlib import Path
    body=[]
    for role,value in blocks(ir):
        if role=='text': body.append('<p>'+html.escape(value).replace('\n','<br>')+'</p>')
        else: body.append('<table>'+''.join('<tr>'+''.join('<td>'+html.escape(str(v or ''))+'</td>' for v in row)+'</tr>' for row in value)+'</table>')
    with tempfile.TemporaryDirectory(prefix='transformx-reflow-') as folder:
        for i,image in enumerate(images):
            Path(folder,f'asset-{i}.png').write_bytes(image)
            body.append(f'<p><img src="asset-{i}.png" width="450"></p>')
        story=pymupdf.Story(html='<html><body>'+''.join(body)+'</body></html>',user_css='body{font-size:11pt} table{border-collapse:collapse;width:100%}td{border:1px solid #aaa;padding:5px}',archive=pymupdf.Archive(folder))
        output=BytesIO(); writer=pymupdf.DocumentWriter(output)
        for _ in range(500):
            device=writer.begin_page(pymupdf.Rect(0,0,595,842))
            more,_=story.place(pymupdf.Rect(36,36,559,806)); story.draw(device); writer.end_page()
            if not more: break
        else:
            writer.close(); raise ValueError('Reflow output exceeded 500 pages.')
        writer.close()
        return output.getvalue()
