"""Fixed-page Office export: artwork background plus editable native text.

PDF vector artwork is rasterized at 150 dpi. Text and table cell coordinates
remain page-relative. This preserves appearance, not original drawing objects.
"""
from io import BytesIO
import pymupdf
from extraction.coordinator import text_regions


def pages(data,ir):
    with pymupdf.open(stream=data,filetype='pdf') as source:
        for unit in ir['units']:
            page=source[unit['number']-1]
            regions=[r for r in text_regions({'units':[unit]}) if r.get('bbox')]
            spans=[s for b in page.get_text('dict')['blocks'] for line in b.get('lines',[]) for s in line['spans']]
            for region in regions:
                matches=[s for s in spans if s['text'].strip()==region['text'].strip()]
                if len(matches)==1:
                    region['_export_font_size']=matches[0]['size']
                    region['_export_font']=matches[0]['font']
            # Exact glyph removal retains images and vector artwork. We use
            # native span boxes, not uncertain semantic region rectangles.
            for block in page.get_text('dict')['blocks']:
                for line in block.get('lines',[]):
                    for span in line['spans']: page.add_redact_annot(span['bbox'],fill=False)
            page.apply_redactions(images=0,graphics=0,text=0)
            yield unit,page.get_pixmap(matrix=pymupdf.Matrix(150/72,150/72),alpha=False).tobytes('png'),regions


def export(data,ir,kind):
    output=BytesIO()
    records=list(pages(data,ir))
    if kind=='pptx':
        from pptx import Presentation
        from pptx.util import Pt
        from pptx.enum.text import MSO_AUTO_SIZE
        deck=Presentation()
        if records: deck.slide_width=Pt(records[0][0]['width']); deck.slide_height=Pt(records[0][0]['height'])
        for unit,image,regions in records:
            slide=deck.slides.add_slide(deck.slide_layouts[6])
            sx=deck.slide_width/Pt(unit['width']); sy=deck.slide_height/Pt(unit['height'])
            slide.shapes.add_picture(BytesIO(image),0,0,width=deck.slide_width,height=deck.slide_height)
            for region in regions:
                x0,y0,x1,y1=region['bbox']; text=region.get('transformed_text',region['text'])
                shape=slide.shapes.add_textbox(int(Pt(x0)*sx),int(Pt(y0)*sy),int(Pt(x1-x0)*sx),int(Pt(y1-y0)*sy))
                shape.name=region['id']; frame=shape.text_frame
                frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
                frame.word_wrap=False; frame.auto_size=MSO_AUTO_SIZE.NONE
                frame.text=text
                for paragraph in frame.paragraphs:
                    paragraph.space_before=paragraph.space_after=Pt(0); paragraph.line_spacing=1.0
                    paragraph.font.name='Malgun Gothic' if any(ord(c)>127 for c in text) else 'Arial'
                    paragraph.font.size=Pt(region.get('_export_font_size',max(8,min(14,(y1-y0)/max(1,len(text.splitlines()))*.65))))
        deck.save(output)
    elif kind=='docx':
        from docx import Document
        from docx.shared import Pt
        from docx.oxml import parse_xml
        from docx.enum.section import WD_SECTION_START
        from xml.sax.saxutils import escape
        doc=Document()
        for index,(unit,image,regions) in enumerate(records):
            section=doc.sections[0] if index==0 else doc.add_section(WD_SECTION_START.NEW_PAGE)
            section.page_width=Pt(unit['width']); section.page_height=Pt(unit['height'])
            section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Pt(0)
            # Background attached to this section's header avoids flow pushing
            # a full-page image onto an extra page. Each section is independent.
            section.header.is_linked_to_previous=False
            header=section.header.paragraphs[0]; run=header.add_run()
            inline=run.add_picture(BytesIO(image),width=Pt(unit['width']),height=Pt(unit['height']))._inline
            anchor=parse_xml('<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH><wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV><wp:wrapNone/></wp:anchor>')
            for child in list(inline): anchor.append(child)
            wrap=anchor.find('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}wrapNone')
            anchor.remove(wrap)
            extent=anchor.find('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent')
            anchor.insert(list(anchor).index(extent)+1,wrap)
            inline.getparent().replace(inline,anchor)
            paragraph=doc.add_paragraph(); paragraph.paragraph_format.space_after=Pt(0)
            for region in regions:
                x0,y0,x1,y1=region['bbox']; text=region.get('transformed_text',region['text'])
                size=max(8,region.get('_export_font_size',min(14,(y1-y0)/max(1,len(text.splitlines()))*.65))*.85)
                xml='<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><v:shape type="#_x0000_t202" style="position:absolute;margin-left:'+str(x0)+'pt;margin-top:'+str(y0)+'pt;width:'+str(x1-x0)+'pt;height:'+str(y1-y0)+'pt;mso-position-horizontal-relative:page;mso-position-vertical-relative:page" stroked="f" filled="f"><v:textbox inset="0,0,0,0"><w:txbxContent><w:p><w:pPr><w:spacing w:before="0" w:after="0" w:line="'+str(round(size*20))+'" w:lineRule="exact"/></w:pPr><w:r><w:rPr><w:rFonts w:ascii="Arial" w:eastAsia="Malgun Gothic"/><w:sz w:val="'+str(round(size*2))+'"/></w:rPr><w:t xml:space="preserve">'+escape(text)+'</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict>'
                paragraph.add_run()._r.append(parse_xml(xml))
        doc.save(output)
    else: raise ValueError('Layout export supports DOCX and PPTX.')
    return output.getvalue()
